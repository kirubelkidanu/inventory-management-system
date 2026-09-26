# backend/app/services/report.py
from datetime import date, datetime, timezone
from decimal import Decimal
import io
from typing import Any, Dict, List, Optional
import uuid

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.repositories.report import ReportRepository
from app.schemas.report import (
    InTransitReportItem,
    StockBalanceItem,
    StockBalanceResponse,
    TrialBalanceItem,
    TrialBalanceResponse,
)


class ReportService:
    def __init__(self, session: AsyncSession):
        self.session = session
        self.repo = ReportRepository(session)

    def _resolve_reference_number(self, row: Dict[str, Any]) -> Optional[str]:
        """Extracts the best document reference based on available transaction fields."""
        return (
            row.get("invoice_no")
            or row.get("siv_no")
            or row.get("istv_no")
            or row.get("received_grv_no")
            or row.get("requested_no")
            or row.get("external_reference")
            or row.get("adjustment_reason")
        )

    def _map_trial_balance_item(self, row: Dict[str, Any]) -> TrialBalanceItem:
        """Converts raw query dictionary into a validated TrialBalanceItem."""
        movement_type = row.get("movement_type", "IN")
        quantity = Decimal(str(row.get("quantity", "0.0000")))
        if movement_type == "IN":
            in_qty = quantity
            out_qty = Decimal("0.0000")
        else:
            in_qty = Decimal("0.0000")
            out_qty = quantity

        return TrialBalanceItem(
            movement_id=row["movement_id"],
            movement_date=row["movement_date"],
            created_at=row["created_at"],
            transaction_id=row["transaction_id"],
            transaction_number=row["transaction_number"],
            transaction_type=row["transaction_type"],
            reference_number=self._resolve_reference_number(row),
            item_id=row["item_id"],
            item_code=row["item_code"],
            item_description=row["item_description"],
            category_name=row["category_name"],
            unit=row["unit"],
            warehouse_id=row["warehouse_id"],
            warehouse_code=row["warehouse_code"],
            warehouse_name=row["warehouse_name"],
            project_name=row.get("project_name"),
            plate_no=row.get("plate_no"),
            driver_name=row.get("driver_name"),
            in_quantity=in_qty,
            out_quantity=out_qty,
            signed_quantity=Decimal(str(row.get("signed_quantity", "0.0000"))),
            running_balance=(
                Decimal(str(row["running_balance"]))
                if row.get("running_balance") is not None
                else None
            ),
            status=row.get("status", "POSTED"),
            entered_by_name=row.get("entered_by_name", "System"),
        )

    async def get_trial_balance_report(
        self,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
        warehouse_id: Optional[uuid.UUID] = None,
        project_id: Optional[uuid.UUID] = None,
        item_id: Optional[uuid.UUID] = None,
        category_id: Optional[uuid.UUID] = None,
        transaction_type: Optional[str] = None,
        search: Optional[str] = None,
        page: int = 1,
        page_size: int = 50,
    ) -> TrialBalanceResponse:
        """Generates paginated Trial Balance response with aggregate sums."""
        page = max(1, page)
        page_size = max(1, min(page_size, 500))
        offset = (page - 1) * page_size

        rows, total_count, total_in, total_out = await self.repo.get_trial_balance(
            date_from=date_from,
            date_to=date_to,
            warehouse_id=warehouse_id,
            project_id=project_id,
            item_id=item_id,
            category_id=category_id,
            transaction_type=transaction_type,
            search=search,
            limit=page_size,
            offset=offset,
        )

        items = [self._map_trial_balance_item(r) for r in rows]
        return TrialBalanceResponse(
            total_count=total_count,
            page=page,
            page_size=page_size,
            items=items,
            total_in=total_in,
            total_out=total_out,
        )

    async def export_trial_balance_excel(
        self,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
        warehouse_id: Optional[uuid.UUID] = None,
        project_id: Optional[uuid.UUID] = None,
        item_id: Optional[uuid.UUID] = None,
        category_id: Optional[uuid.UUID] = None,
        transaction_type: Optional[str] = None,
        search: Optional[str] = None,
    ) -> io.BytesIO:
        """
        Exports the entire filtered trial balance to an Excel spreadsheet with styled
        headers, auto-adjusted column widths, and aggregate formulas/totals.
        """
        raw_rows = await self.repo.get_all_trial_balance_movements(
            date_from=date_from,
            date_to=date_to,
            warehouse_id=warehouse_id,
            project_id=project_id,
            item_id=item_id,
            category_id=category_id,
            transaction_type=transaction_type,
            search=search,
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Trial Balance"

        # Define Styles
        title_font = Font(name="Calibri", size=16, bold=True, color="1F497D")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        total_font = Font(name="Calibri", size=11, bold=True)
        total_fill = PatternFill(start_color="E9EDF4", end_color="E9EDF4", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        double_bottom_border = Border(
            top=Side(style="thin", color="000000"),
            bottom=Side(style="double", color="000000"),
        )

        # Title Block
        ws.merge_cells("A1:R1")
        title_cell = ws["A1"]
        title_cell.value = "DETAILED INVENTORY MOVEMENT / TRIAL BALANCE"
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # Metadata Row
        generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        ws.merge_cells("A2:R2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Generated: {generated_at} | Total Records: {len(raw_rows)}"
        sub_cell.font = Font(name="Calibri", size=10, italic=True, color="595959")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 18

        # Headers
        headers = [
            "Date",
            "Tx Number",
            "Type",
            "Ref / Doc No",
            "Item Code",
            "Description",
            "Category",
            "Unit",
            "WH Code",
            "Warehouse Name",
            "Project",
            "Plate No",
            "Driver Name",
            "IN",
            "OUT",
            "Running Bal",
            "Status",
            "Entered By",
        ]

        ws.append([])  # blank row 3
        ws.append(headers)  # row 4
        ws.row_dimensions[4].height = 24

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Data Rows
        sum_in = Decimal("0.0000")
        sum_out = Decimal("0.0000")

        for row_dict in raw_rows:
            item = self._map_trial_balance_item(row_dict)
            sum_in += item.in_quantity
            sum_out += item.out_quantity

            row_data = [
                item.movement_date.strftime("%Y-%m-%d"),
                item.transaction_number,
                item.transaction_type,
                item.reference_number or "",
                item.item_code,
                item.item_description,
                item.category_name,
                item.unit,
                item.warehouse_code,
                item.warehouse_name,
                item.project_name or "",
                item.plate_no or "",
                item.driver_name or "",
                float(item.in_quantity),
                float(item.out_quantity),
                float(item.running_balance) if item.running_balance is not None else "",
                item.status,
                item.entered_by_name,
            ]
            ws.append(row_data)
            current_row = ws.max_row
            ws.row_dimensions[current_row].height = 19

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                # Align numbers and format
                if col_idx in (14, 15, 16):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0.0000"
                elif col_idx in (1, 3, 8, 9, 17):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Total Row
        total_row_idx = ws.max_row + 1
        ws.append(
            [
                "TOTAL",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                float(sum_in),
                float(sum_out),
                "",
                "",
                "",
            ]
        )
        ws.row_dimensions[total_row_idx].height = 22
        ws.merge_cells(f"A{total_row_idx}:M{total_row_idx}")
        total_label = ws[f"A{total_row_idx}"]
        total_label.alignment = Alignment(horizontal="right", vertical="center")

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=total_row_idx, column=col_idx)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = double_bottom_border
            if col_idx in (14, 15):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.0000"

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in (1, 2):  # skip merged header lines
                    continue
                val = cell.value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    async def get_stock_balance_report(
        self,
        warehouse_id: Optional[uuid.UUID] = None,
        category_id: Optional[uuid.UUID] = None,
        search: Optional[str] = None,
    ) -> StockBalanceResponse:
        """Returns current stock balances with calculated available quantities."""
        rows = await self.repo.get_current_stock_balances(
            warehouse_id=warehouse_id,
            category_id=category_id,
            search=search,
        )
        items = [
            StockBalanceItem(
                warehouse_id=r["warehouse_id"],
                warehouse_code=r["warehouse_code"],
                warehouse_name=r["warehouse_name"],
                project_name=r.get("project_name"),
                category_name=r["category_name"],
                item_id=r["item_id"],
                item_code=r["item_code"],
                item_description=r["item_description"],
                unit=r["unit"],
                quantity_on_hand=Decimal(str(r["quantity_on_hand"])),
                quantity_reserved=Decimal(str(r["quantity_reserved"])),
                quantity_available=Decimal(str(r["quantity_available"])),
            )
            for r in rows
        ]
        return StockBalanceResponse(total_count=len(items), items=items)

    async def get_in_transit_report(
        self,
        source_warehouse_id: Optional[uuid.UUID] = None,
        destination_warehouse_id: Optional[uuid.UUID] = None,
        status: Optional[str] = None,
    ) -> List[InTransitReportItem]:
        """Returns transfer items currently in transit with remaining quantities."""
        rows = await self.repo.get_in_transit_transfers(
            source_warehouse_id=source_warehouse_id,
            destination_warehouse_id=destination_warehouse_id,
            status=status,
        )
        return [
            InTransitReportItem(
                transfer_record_id=r["transfer_record_id"],
                istv_transaction_id=r["istv_transaction_id"],
                istv_number=r["istv_number"],
                transaction_date=r["transaction_date"],
                source_warehouse_name=r["source_warehouse_name"],
                destination_warehouse_name=r["destination_warehouse_name"],
                plate_no=r.get("plate_no"),
                driver_name=r.get("driver_name"),
                status=r["status"],
                item_id=r["item_id"],
                item_code=r["item_code"],
                item_description=r["item_description"],
                unit=r["unit"],
                sent_quantity=Decimal(str(r["sent_quantity"])),
                received_quantity=Decimal(str(r["received_quantity"])),
                remaining_quantity=Decimal(str(r["remaining_quantity"])),
            )
            for r in rows
        ]
