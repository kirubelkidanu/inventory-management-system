# backend/tests/test_e2e_reconciliation.py
"""
End-to-End System Verification & Reconciliation Audit Test Suite

Verifies the complete inventory lifecycle across two warehouses:
1. Setup: WH001 ("Main Store"), WH002 ("Site Store"), Project, Category, 2 Master Items.
2. Step 1: GRV Receipt (+500 Cement, +200 Rebar into WH001).
3. Step 2: SIV Issue (-120 Cement from WH001 with 4 signatories).
4. Step 3: ISTV Transfer Dispatch (-100 Cement from WH001 to WH002, in-transit).
5. Step 4: ISTRV Partial Receipt (+60 Cement at WH002, status PARTIALLY_RECEIVED).
6. Step 5: ISTRV Final Receipt (+40 Cement at WH002, status COMPLETED).
7. Step 6: SRV Store Return (+15 Cement returned to WH001 referencing original SIV).
8. Step 7: Audit Count & Adjustment (WH001 deficit -3 Cement, surplus +5 Rebar).
9. Step 8: Database Reconciliation Invariant Check (vw_reconciliation_discrepancies == 0).
10. Step 9: Authoritative Trial Balance Ledger Audit (Total IN = 820, OUT = 223, System = 597).
11. Step 10: Excel Export Integrity Verification (binary openpyxl parsing).
"""

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
import io
from typing import Any, Dict, List, Optional
from unittest.mock import AsyncMock, MagicMock
from uuid import UUID, uuid4

from fastapi import HTTPException
from httpx import ASGITransport, AsyncClient
from jose import jwt
import openpyxl
import pytest

from app.api.routes.reconciliation import get_reconciliation_service
from app.api.routes.reports import get_report_service
from app.api.routes.transactions import get_transaction_service
from app.core.config import settings
from app.main import app
from app.models.audit import AuditLog
from app.models.category import Category
from app.models.inventory import InventoryBalance
from app.models.item import Item
from app.models.project import Project
from app.models.transaction import (
    StockMovement,
    Transaction,
    TransactionLine,
    TransferLine,
    TransferRecord,
)
from app.models.user import User
from app.models.warehouse import Warehouse
from app.schemas.reconciliation import (
    PhysicalCountItemInput,
    PhysicalCountSubmitRequest,
    ReconciliationCommitRequest,
)
from app.schemas.transaction import (
    AdjustmentCreate,
    AdjustmentLineCreate,
    GRVCreate,
    ISTRVCreate,
    ISTRVLineCreate,
    ISTVCreate,
    SIVCreate,
    SRVCreate,
    SRVLineCreate,
    TransactionLineCreate,
)
from app.services.reconciliation import ReconciliationService, _to_uuid
from app.services.report import ReportService
from app.services.transaction import TransactionService


# ==============================================================================
# 1. Stateful In-Memory System Store & Repository Harness
# ==============================================================================


class StatefulInventoryStore:
    """Authoritative in-memory state simulating PostgreSQL database tables."""

    def __init__(self):
        self.warehouses: Dict[UUID, Warehouse] = {}
        self.projects: Dict[UUID, Project] = {}
        self.categories: Dict[UUID, Category] = {}
        self.items: Dict[UUID, Item] = {}
        self.users: Dict[UUID, User] = {}
        self.balances: Dict[tuple[UUID, UUID], InventoryBalance] = {}  # (wh_id, item_id) -> Balance
        self.transactions: Dict[UUID, Transaction] = {}
        self.transaction_lines: Dict[UUID, TransactionLine] = {}
        self.transfer_records: Dict[UUID, TransferRecord] = {}
        self.transfer_lines: Dict[UUID, TransferLine] = {}
        self.stock_movements: List[StockMovement] = []
        self.audit_logs: List[AuditLog] = []

    def get_balance(self, warehouse_id: UUID, item_id: UUID) -> Decimal:
        bal = self.balances.get((warehouse_id, item_id))
        return bal.quantity_on_hand if bal else Decimal("0.0000")

    def calculate_discrepancies(self) -> List[Dict[str, Any]]:
        """
        Replicates PostgreSQL view `vw_reconciliation_discrepancies`:
        Compares projected inventory_balances.quantity_on_hand against
        the cumulative sum of signed_quantity in immutable stock_movements.
        """
        ledger_sums: Dict[tuple[UUID, UUID], Decimal] = {}
        for sm in self.stock_movements:
            key = (sm.warehouse_id, sm.item_id)
            ledger_sums[key] = ledger_sums.get(key, Decimal("0.0000")) + sm.signed_quantity

        all_keys = set(self.balances.keys()).union(set(ledger_sums.keys()))
        discrepancies = []
        for wh_id, item_id in all_keys:
            bal = self.balances.get((wh_id, item_id))
            projected = bal.quantity_on_hand if bal else Decimal("0.0000")
            ledger = ledger_sums.get((wh_id, item_id), Decimal("0.0000"))
            diff = projected - ledger
            if diff != Decimal("0.0000"):
                wh = self.warehouses.get(wh_id)
                item = self.items.get(item_id)
                discrepancies.append(
                    {
                        "warehouse_id": wh_id,
                        "warehouse_code": wh.code if wh else "",
                        "item_id": item_id,
                        "item_code": item.item_code if item else "",
                        "projected_balance": projected,
                        "ledger_cumulative_balance": ledger,
                        "discrepancy": diff,
                    }
                )
        return discrepancies


class StatefulTransactionRepository:
    """Implements TransactionRepository contracts using the stateful store."""

    def __init__(self, store: StatefulInventoryStore):
        self.store = store

    async def lock_inventory_balances_for_update(
        self, warehouse_id: UUID, item_ids: List[UUID]
    ) -> Dict[UUID, InventoryBalance]:
        result = {}
        for i_id in sorted(item_ids):
            key = (warehouse_id, i_id)
            if key in self.store.balances:
                result[i_id] = self.store.balances[key]
        return result

    async def get_or_create_balance(
        self,
        warehouse_id: UUID,
        item_id: UUID,
        existing_balances: Optional[Dict[UUID, InventoryBalance]] = None,
    ) -> InventoryBalance:
        if existing_balances and item_id in existing_balances:
            return existing_balances[item_id]

        key = (warehouse_id, item_id)
        if key not in self.store.balances:
            bal = InventoryBalance(
                id=uuid4(),
                warehouse_id=warehouse_id,
                item_id=item_id,
                quantity_on_hand=Decimal("0.0000"),
                quantity_reserved=Decimal("0.0000"),
                updated_at=datetime.now(timezone.utc),
            )
            self.store.balances[key] = bal
        return self.store.balances[key]

    async def create_transaction(self, transaction: Transaction) -> Transaction:
        self.store.transactions[transaction.id] = transaction
        return transaction

    async def create_transaction_lines(
        self, lines: List[TransactionLine]
    ) -> List[TransactionLine]:
        for line in lines:
            self.store.transaction_lines[line.id] = line
        return lines

    async def create_stock_movement(
        self,
        transaction_id: UUID,
        transaction_line_id: UUID,
        item_id: UUID,
        warehouse_id: UUID,
        project_id: Optional[UUID],
        movement_type: str,
        quantity: Decimal,
        running_balance: Optional[Decimal],
        movement_date: date,
    ) -> StockMovement:
        abs_qty = abs(quantity)
        signed_qty = abs_qty if movement_type == "IN" else -abs_qty
        sm = StockMovement(
            id=uuid4(),
            transaction_id=transaction_id,
            transaction_line_id=transaction_line_id,
            item_id=item_id,
            warehouse_id=warehouse_id,
            project_id=project_id,
            movement_type=movement_type,
            quantity=abs_qty,
            signed_quantity=signed_qty,
            running_balance=running_balance,
            movement_date=movement_date,
            created_at=datetime.now(timezone.utc),
        )
        self.store.stock_movements.append(sm)
        return sm

    async def create_audit_log(
        self,
        actor_user_id: Optional[UUID],
        action: str,
        entity_type: str,
        entity_id: str,
        before_state: Optional[dict] = None,
        after_state: Optional[dict] = None,
        change_summary: Optional[str] = None,
        reason: Optional[str] = None,
    ) -> AuditLog:
        log = AuditLog(
            id=uuid4(),
            actor_user_id=actor_user_id,
            action=action,
            entity_type=entity_type,
            entity_id=entity_id,
            before_state=before_state,
            after_state=after_state,
            change_summary=change_summary,
            reason=reason,
            created_at=datetime.now(timezone.utc),
        )
        self.store.audit_logs.append(log)
        return log

    async def get_transaction_by_id(self, transaction_id: UUID) -> Optional[Transaction]:
        tx = self.store.transactions.get(transaction_id)
        if tx:
            tx.lines = [
                l for l in self.store.transaction_lines.values() if l.transaction_id == tx.id
            ]
        return tx

    async def create_transfer_record(
        self, transfer_record: TransferRecord
    ) -> TransferRecord:
        self.store.transfer_records[transfer_record.id] = transfer_record
        return transfer_record

    async def create_transfer_lines(
        self, lines: List[TransferLine]
    ) -> List[TransferLine]:
        for l in lines:
            self.store.transfer_lines[l.id] = l
        return lines

    async def get_transfer_record_by_istv_id(
        self, istv_id: UUID
    ) -> Optional[TransferRecord]:
        for tr in self.store.transfer_records.values():
            if tr.istv_transaction_id == istv_id:
                tr.lines = [
                    l
                    for l in self.store.transfer_lines.values()
                    if l.transfer_record_id == tr.id
                ]
                return tr
        return None

    async def get_transfer_record_by_id(
        self, record_id: UUID
    ) -> Optional[TransferRecord]:
        tr = self.store.transfer_records.get(record_id)
        if tr:
            tr.lines = [
                l
                for l in self.store.transfer_lines.values()
                if l.transfer_record_id == tr.id
            ]
        return tr

    async def list_transactions(
        self,
        transaction_type: Optional[str] = None,
        warehouse_id: Optional[UUID] = None,
        search: Optional[str] = None,
        skip: int = 0,
        limit: int = 50,
    ) -> List[Transaction]:
        results = []
        for tx in self.store.transactions.values():
            if transaction_type and tx.transaction_type != transaction_type.upper():
                continue
            if warehouse_id and tx.warehouse_id != warehouse_id:
                continue
            tx.lines = [
                l for l in self.store.transaction_lines.values() if l.transaction_id == tx.id
            ]
            results.append(tx)
        return results[skip : skip + limit]


class StatefulReportRepository:
    """Implements ReportRepository contracts using the stateful store."""

    def __init__(self, store: StatefulInventoryStore):
        self.store = store

    async def get_trial_balance(
        self,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
        warehouse_id: Optional[UUID] = None,
        project_id: Optional[UUID] = None,
        item_id: Optional[UUID] = None,
        category_id: Optional[UUID] = None,
        transaction_type: Optional[str] = None,
        search: Optional[str] = None,
        limit: int = 50,
        offset: int = 0,
    ):
        rows = []
        for sm in self.store.stock_movements:
            tx = self.store.transactions.get(sm.transaction_id)
            line = self.store.transaction_lines.get(sm.transaction_line_id)
            item = self.store.items.get(sm.item_id)
            cat = self.store.categories.get(item.category_id) if item else None
            wh = self.store.warehouses.get(sm.warehouse_id)
            proj = (
                self.store.projects.get(sm.project_id)
                if sm.project_id
                else (self.store.projects.get(tx.project_id) if tx and tx.project_id else None)
            )
            user = self.store.users.get(tx.created_by) if tx and tx.created_by else None

            # Filters
            if date_from and sm.movement_date < date_from:
                continue
            if date_to and sm.movement_date > date_to:
                continue
            if warehouse_id and sm.warehouse_id != warehouse_id:
                continue
            if item_id and sm.item_id != item_id:
                continue
            if category_id and item and item.category_id != category_id:
                continue
            if transaction_type and tx and tx.transaction_type != transaction_type:
                continue

            row = {
                "movement_id": sm.id,
                "movement_date": sm.movement_date,
                "created_at": sm.created_at,
                "transaction_id": sm.transaction_id,
                "transaction_number": tx.transaction_number if tx else "UNKNOWN",
                "transaction_type": tx.transaction_type if tx else "UNKNOWN",
                "invoice_no": getattr(tx, "invoice_no", None),
                "siv_no": getattr(tx, "siv_no", None),
                "istv_no": getattr(tx, "istv_no", None),
                "received_grv_no": getattr(tx, "received_grv_no", None),
                "requested_no": getattr(tx, "requested_no", None),
                "external_reference": None,
                "adjustment_reason": getattr(tx, "adjustment_reason", None),
                "remarks": getattr(tx, "remarks", None),
                "item_id": sm.item_id,
                "item_code": item.item_code if item else "",
                "item_description": item.description if item else "",
                "category_name": cat.name if cat else "",
                "unit": (line.unit if line else (item.default_unit if item else "PCS")),
                "warehouse_id": sm.warehouse_id,
                "warehouse_code": wh.code if wh else "",
                "warehouse_name": wh.name if wh else "",
                "project_name": proj.name if proj else None,
                "plate_no": getattr(tx, "plate_no", None),
                "driver_name": getattr(tx, "driver_name", None),
                "movement_type": sm.movement_type,
                "quantity": sm.quantity,
                "signed_quantity": sm.signed_quantity,
                "running_balance": sm.running_balance,
                "status": tx.status if tx else "POSTED",
                "entered_by_name": user.full_name if user else "Store Clerk",
            }
            rows.append(row)

        total_count = len(rows)
        total_in = sum(r["quantity"] for r in rows if r["movement_type"] == "IN")
        total_out = sum(r["quantity"] for r in rows if r["movement_type"] == "OUT")
        paginated_rows = rows[offset : offset + limit]

        return paginated_rows, total_count, total_in, total_out

    async def get_all_trial_balance_movements(
        self,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
        warehouse_id: Optional[UUID] = None,
        project_id: Optional[UUID] = None,
        item_id: Optional[UUID] = None,
        category_id: Optional[UUID] = None,
        transaction_type: Optional[str] = None,
        search: Optional[str] = None,
    ) -> List[Dict[str, Any]]:
        rows, _, _, _ = await self.get_trial_balance(
            date_from=date_from,
            date_to=date_to,
            warehouse_id=warehouse_id,
            project_id=project_id,
            item_id=item_id,
            category_id=category_id,
            transaction_type=transaction_type,
            search=search,
            limit=10000,
            offset=0,
        )
        return rows


def create_stateful_session_mock(store: StatefulInventoryStore):
    """Mocks AsyncSession.execute for queries made inside ReconciliationService."""
    mock_session = AsyncMock()

    async def fake_execute(stmt):
        stmt_str = str(stmt)
        # 1. Warehouse query
        if "FROM warehouses" in stmt_str or "warehouses.id" in stmt_str:
            mock_res = MagicMock()
            for wh in store.warehouses.values():
                if str(wh.id) in stmt_str or hasattr(stmt, "compile"):
                    mock_res.scalar_one_or_none.return_value = wh
                    return mock_res
            mock_res.scalar_one_or_none.return_value = next(
                iter(store.warehouses.values()), None
            )
            return mock_res

        # 2. Items query
        if "FROM items" in stmt_str or "items.id" in stmt_str:
            mock_res = MagicMock()
            mock_scalars = MagicMock()
            mock_scalars.all.return_value = list(store.items.values())
            mock_res.scalars.return_value = mock_scalars
            return mock_res

        # 3. Default fallback
        mock_res = MagicMock()
        mock_res.scalar_one_or_none.return_value = None
        mock_res.scalars.return_value.all.return_value = []
        return mock_res

    mock_session.execute = AsyncMock(side_effect=fake_execute)
    mock_session.commit = AsyncMock()
    mock_session.rollback = AsyncMock()
    return mock_session


# ==============================================================================
# 2. Fixtures & Test Setup
# ==============================================================================


@pytest.fixture
def token_factory():
    def make_token(role: str = "ADMIN", user_id: str = "user-123") -> str:
        return jwt.encode(
            {
                "sub": user_id,
                "email": f"{user_id}@example.com",
                "role": role,
                "exp": datetime.now(timezone.utc) + timedelta(minutes=30),
            },
            settings.JWT_SECRET,
            algorithm=settings.JWT_ALGORITHM,
        )

    return make_token


@pytest.fixture
def e2e_system():
    """Builds the complete stateful inventory system test harness."""
    store = StatefulInventoryStore()

    # Create Master Entities
    wh001 = Warehouse(
        id=uuid4(),
        code="WH001",
        name="Main Store",
        location="Headquarters Central Yard",
        is_active=True,
        created_at=datetime.now(timezone.utc),
    )
    wh002 = Warehouse(
        id=uuid4(),
        code="WH002",
        name="Site Store",
        location="Kazanchis Project Depot",
        is_active=True,
        created_at=datetime.now(timezone.utc),
    )
    store.warehouses[wh001.id] = wh001
    store.warehouses[wh002.id] = wh002

    project = Project(
        id=uuid4(),
        code="PRJ001",
        name="Kazanchis Commercial Complex",
        location="Kazanchis, Addis Ababa",
        is_active=True,
        created_at=datetime.now(timezone.utc),
    )
    store.projects[project.id] = project

    category = Category(
        id=uuid4(),
        code="CON",
        name="Construction Materials",
        description="Structural and building materials",
        is_active=True,
        created_at=datetime.now(timezone.utc),
    )
    store.categories[category.id] = category

    cement = Item(
        id=uuid4(),
        item_code="01-CM-00001",
        description="Portland Cement 42.5N",
        category_id=category.id,
        default_unit="BAG",
        is_active=True,
        created_at=datetime.now(timezone.utc),
    )
    rebar = Item(
        id=uuid4(),
        item_code="02-ST-00001",
        description="Reinforcement Bar 12mm",
        category_id=category.id,
        default_unit="PCS",
        is_active=True,
        created_at=datetime.now(timezone.utc),
    )
    store.items[cement.id] = cement
    store.items[rebar.id] = rebar

    admin_user = User(
        id=uuid4(),
        email="admin@inventory.corp",
        full_name="Abebe Store Manager",
        role="ADMIN",
        is_active=True,
        created_at=datetime.now(timezone.utc),
    )
    store.users[admin_user.id] = admin_user

    # Wire stateful services
    mock_session = create_stateful_session_mock(store)

    tx_repo = StatefulTransactionRepository(store)
    tx_service = TransactionService(mock_session)
    tx_service.repo = tx_repo

    reconciliation_service = ReconciliationService(mock_session)
    reconciliation_service.tx_service = tx_service
    reconciliation_service.tx_repo = tx_repo

    report_repo = StatefulReportRepository(store)
    report_service = ReportService(mock_session)
    report_service.repo = report_repo

    return {
        "store": store,
        "wh001": wh001,
        "wh002": wh002,
        "project": project,
        "category": category,
        "cement": cement,
        "rebar": rebar,
        "user": admin_user,
        "tx_service": tx_service,
        "reconciliation_service": reconciliation_service,
        "report_service": report_service,
    }


# ==============================================================================
# 3. End-to-End Comprehensive Lifecycle & Reconciliation Audit
# ==============================================================================


@pytest.mark.asyncio
async def test_full_lifecycle_e2e_reconciliation_and_trial_balance(e2e_system):
    """
    Executes the entire 10-step lifecycle and verifies all invariants:
    Dual-write stock updates, transfer logistics, store returns, physical reconciliation,
    database view zero-discrepancy invariant, and Trial Balance ledger integrity.
    """
    store: StatefulInventoryStore = e2e_system["store"]
    wh001: Warehouse = e2e_system["wh001"]
    wh002: Warehouse = e2e_system["wh002"]
    project: Project = e2e_system["project"]
    cement: Item = e2e_system["cement"]
    rebar: Item = e2e_system["rebar"]
    user: User = e2e_system["user"]
    tx_service: TransactionService = e2e_system["tx_service"]
    reconciliation_service: ReconciliationService = e2e_system["reconciliation_service"]
    report_service: ReportService = e2e_system["report_service"]

    # --------------------------------------------------------------------------
    # Step 1: GRV Receipt (+500 Cement, +200 Rebar into WH001)
    # --------------------------------------------------------------------------
    grv_payload = GRVCreate(
        warehouse_id=wh001.id,
        project_id=project.id,
        supplier_name="National Cement Share Co.",
        invoice_no="INV-2026-0881",
        received_grv_no="GRV-EXT-4401",
        remarks="Initial project foundation batch receipt",
        lines=[
            TransactionLineCreate(
                item_id=cement.id, quantity=Decimal("500.0000"), unit="BAG"
            ),
            TransactionLineCreate(
                item_id=rebar.id, quantity=Decimal("200.0000"), unit="PCS"
            ),
        ],
    )
    grv_tx = await tx_service.create_and_post_grv(user_id=user.id, payload=grv_payload)

    assert grv_tx.transaction_type == "GRV"
    assert grv_tx.status == "POSTED"
    assert store.get_balance(wh001.id, cement.id) == Decimal("500.0000")
    assert store.get_balance(wh001.id, rebar.id) == Decimal("200.0000")

    grv_movements = [sm for sm in store.stock_movements if sm.transaction_id == grv_tx.id]
    assert len(grv_movements) == 2
    assert all(sm.movement_type == "IN" and sm.signed_quantity > 0 for sm in grv_movements)

    # --------------------------------------------------------------------------
    # Step 2: SIV Issue (-120 Cement from WH001 with 4 Signatories)
    # --------------------------------------------------------------------------
    siv_payload = SIVCreate(
        warehouse_id=wh001.id,
        project_id=project.id,
        requested_from="Civil Works Section",
        project_dept="Structural Engineering",
        requested_no="REQ-2026-019",
        issued_by_name="Almaz Storekeeper",
        checked_by_name="Dawit Senior Inspector",
        received_by_name="Chala Site Foreman",
        approved_by_name="Bekele Project Manager",
        remarks="Issued for Foundation Slab Pouring Zone A",
        lines=[
            TransactionLineCreate(
                item_id=cement.id, quantity=Decimal("120.0000"), unit="BAG"
            )
        ],
    )
    siv_tx = await tx_service.create_and_post_siv(user_id=user.id, payload=siv_payload)

    assert siv_tx.transaction_type == "SIV"
    assert siv_tx.status == "POSTED"
    assert siv_tx.issued_by_name == "Almaz Storekeeper"
    assert siv_tx.checked_by_name == "Dawit Senior Inspector"
    assert siv_tx.received_by_name == "Chala Site Foreman"
    assert siv_tx.approved_by_name == "Bekele Project Manager"
    assert store.get_balance(wh001.id, cement.id) == Decimal("380.0000")

    siv_mov = next(sm for sm in store.stock_movements if sm.transaction_id == siv_tx.id)
    assert siv_mov.movement_type == "OUT"
    assert siv_mov.signed_quantity == Decimal("-120.0000")
    assert siv_mov.running_balance == Decimal("380.0000")

    # --------------------------------------------------------------------------
    # Step 3: ISTV Transfer Dispatch (-100 Cement from WH001 to WH002)
    # --------------------------------------------------------------------------
    istv_payload = ISTVCreate(
        source_warehouse_id=wh001.id,
        destination_warehouse_id=wh002.id,
        project_id=project.id,
        plate_no="3-A12345",
        driver_name="Abebe",
        remarks="Urgent transfer of cement to site satellite store",
        lines=[
            TransactionLineCreate(
                item_id=cement.id, quantity=Decimal("100.0000"), unit="BAG"
            )
        ],
    )
    istv_tx = await tx_service.create_and_post_istv(user_id=user.id, payload=istv_payload)

    assert istv_tx.transaction_type == "ISTV"
    assert istv_tx.status == "POSTED"
    assert istv_tx.plate_no == "3-A12345"
    assert istv_tx.driver_name == "Abebe"

    # Source decremented, destination UNCHANGED while in transit
    assert store.get_balance(wh001.id, cement.id) == Decimal("280.0000")
    assert store.get_balance(wh002.id, cement.id) == Decimal("0.0000")

    # Transfer record state verified
    tr = await tx_service.get_transfer_record_by_istv_id(istv_tx.id)
    assert tr is not None
    assert tr.status == "IN_TRANSIT"
    tl = tr.lines[0]
    assert tl.sent_quantity == Decimal("100.0000")
    assert tl.received_quantity == Decimal("0.0000")
    assert (tl.sent_quantity - tl.received_quantity) == Decimal("100.0000")

    # --------------------------------------------------------------------------
    # Step 4: ISTRV Partial Receipt (60 Cement received at WH002)
    # --------------------------------------------------------------------------
    istrv1_payload = ISTRVCreate(
        istv_transaction_id=istv_tx.id,
        destination_warehouse_id=wh002.id,
        remarks="Partial shipment arrival - 60 bags received",
        lines=[ISTRVLineCreate(item_id=cement.id, quantity=Decimal("60.0000"))],
    )
    istrv1_tx = await tx_service.create_and_post_istrv(
        user_id=user.id, payload=istrv1_payload
    )

    assert istrv1_tx.transaction_type == "ISTRV"
    assert istrv1_tx.status == "POSTED"
    assert store.get_balance(wh002.id, cement.id) == Decimal("60.0000")

    tr = await tx_service.get_transfer_record_by_istv_id(istv_tx.id)
    assert tr.status == "PARTIALLY_RECEIVED"
    tl = tr.lines[0]
    assert tl.received_quantity == Decimal("60.0000")
    assert (tl.sent_quantity - tl.received_quantity) == Decimal("40.0000")

    # --------------------------------------------------------------------------
    # Step 5: ISTRV Final Receipt (Remaining 40 Cement received at WH002)
    # --------------------------------------------------------------------------
    istrv2_payload = ISTRVCreate(
        istv_transaction_id=istv_tx.id,
        destination_warehouse_id=wh002.id,
        remarks="Final shipment arrival - remaining 40 bags received",
        lines=[ISTRVLineCreate(item_id=cement.id, quantity=Decimal("40.0000"))],
    )
    istrv2_tx = await tx_service.create_and_post_istrv(
        user_id=user.id, payload=istrv2_payload
    )

    assert istrv2_tx.transaction_type == "ISTRV"
    assert istrv2_tx.status == "POSTED"
    assert store.get_balance(wh002.id, cement.id) == Decimal("100.0000")

    tr = await tx_service.get_transfer_record_by_istv_id(istv_tx.id)
    assert tr.status == "COMPLETED"
    assert tr.completed_at is not None
    tl = tr.lines[0]
    assert tl.received_quantity == Decimal("100.0000")
    assert (tl.sent_quantity - tl.received_quantity) == Decimal("0.0000")

    # --------------------------------------------------------------------------
    # Step 6: SRV Store Return (+15 Cement returned to WH001 referencing SIV)
    # --------------------------------------------------------------------------
    srv_payload = SRVCreate(
        reference_siv_id=siv_tx.id,
        warehouse_id=wh001.id,
        project_id=project.id,
        remarks="Unused cement bags returned from Zone A pouring",
        lines=[
            SRVLineCreate(item_id=cement.id, quantity=Decimal("15.0000"), unit="BAG")
        ],
    )
    srv_tx = await tx_service.create_and_post_srv(user_id=user.id, payload=srv_payload)

    assert srv_tx.transaction_type == "SRV"
    assert srv_tx.status == "POSTED"
    assert srv_tx.reference_transaction_id == siv_tx.id
    assert store.get_balance(wh001.id, cement.id) == Decimal("295.0000")

    srv_mov = next(sm for sm in store.stock_movements if sm.transaction_id == srv_tx.id)
    assert srv_mov.movement_type == "IN"
    assert srv_mov.signed_quantity == Decimal("15.0000")
    assert srv_mov.running_balance == Decimal("295.0000")

    # --------------------------------------------------------------------------
    # Step 7: Audit Count & Adjustment (WH001: Cement counted=292, Rebar counted=205)
    # --------------------------------------------------------------------------
    reconciliation_payload = ReconciliationCommitRequest(
        warehouse_id=wh001.id,
        count_date=date.today(),
        count_reference="COUNT-20260919-WH001",
        adjustment_reason="Quarterly Physical Audit Variance",
        counts=[
            PhysicalCountItemInput(
                item_id=cement.id, counted_quantity=Decimal("292.0000")
            ),
            PhysicalCountItemInput(
                item_id=rebar.id, counted_quantity=Decimal("205.0000")
            ),
        ],
    )
    rec_result = await reconciliation_service.commit_reconciliation_adjustments(
        user_id=user.id, payload=reconciliation_payload
    )

    assert rec_result["count_reference"] == "COUNT-20260919-WH001"
    assert rec_result["adjusted_lines_count"] == 2
    adj_tx_number = rec_result["transaction_number"]
    assert adj_tx_number.startswith("ADJ-")

    # Assert final balances
    assert store.get_balance(wh001.id, cement.id) == Decimal("292.0000")
    assert store.get_balance(wh001.id, rebar.id) == Decimal("205.0000")
    assert store.get_balance(wh002.id, cement.id) == Decimal("100.0000")

    # Compensating movements checked
    adj_tx = next(
        tx for tx in store.transactions.values() if tx.transaction_number == adj_tx_number
    )
    adj_movements = [sm for sm in store.stock_movements if sm.transaction_id == adj_tx.id]
    assert len(adj_movements) == 2

    cement_adj_mov = next(sm for sm in adj_movements if sm.item_id == cement.id)
    assert cement_adj_mov.movement_type == "OUT"
    assert cement_adj_mov.signed_quantity == Decimal("-3.0000")
    assert cement_adj_mov.running_balance == Decimal("292.0000")

    rebar_adj_mov = next(sm for sm in adj_movements if sm.item_id == rebar.id)
    assert rebar_adj_mov.movement_type == "IN"
    assert rebar_adj_mov.signed_quantity == Decimal("5.0000")
    assert rebar_adj_mov.running_balance == Decimal("205.0000")

    # --------------------------------------------------------------------------
    # Step 8: Database Reconciliation Invariant Check (vw_reconciliation_discrepancies)
    # --------------------------------------------------------------------------
    discrepancies = store.calculate_discrepancies()
    assert (
        len(discrepancies) == 0
    ), f"Expected 0 reconciliation discrepancies, but found: {discrepancies}"

    # Detailed item-by-item invariant assertions
    # WH001 Cement: +500 (GRV) - 120 (SIV) - 100 (ISTV) + 15 (SRV) - 3 (ADJ) = 292
    wh001_cement_sum = sum(
        sm.signed_quantity
        for sm in store.stock_movements
        if sm.warehouse_id == wh001.id and sm.item_id == cement.id
    )
    assert wh001_cement_sum == Decimal("292.0000")
    assert store.get_balance(wh001.id, cement.id) == wh001_cement_sum

    # WH001 Rebar: +200 (GRV) + 5 (ADJ) = 205
    wh001_rebar_sum = sum(
        sm.signed_quantity
        for sm in store.stock_movements
        if sm.warehouse_id == wh001.id and sm.item_id == rebar.id
    )
    assert wh001_rebar_sum == Decimal("205.0000")
    assert store.get_balance(wh001.id, rebar.id) == wh001_rebar_sum

    # WH002 Cement: +60 (ISTRV) + 40 (ISTRV) = 100
    wh002_cement_sum = sum(
        sm.signed_quantity
        for sm in store.stock_movements
        if sm.warehouse_id == wh002.id and sm.item_id == cement.id
    )
    assert wh002_cement_sum == Decimal("100.0000")
    assert store.get_balance(wh002.id, cement.id) == wh002_cement_sum

    # --------------------------------------------------------------------------
    # Step 9: Authoritative Trial Balance Ledger Audit
    # --------------------------------------------------------------------------
    tb_report = await report_service.get_trial_balance_report(page=1, page_size=100)

    # 9 total movements:
    # 2 (GRV) + 1 (SIV) + 1 (ISTV) + 1 (ISTRV1) + 1 (ISTRV2) + 1 (SRV) + 2 (ADJ) = 9
    assert tb_report.total_count == 9
    assert len(tb_report.items) == 9

    # Cumulative Gross Totals:
    # IN: 500 + 200 + 60 + 40 + 15 + 5 = 820
    assert tb_report.total_in == Decimal("820.0000")
    # OUT: 120 + 100 + 3 = 223
    assert tb_report.total_out == Decimal("223.0000")
    # Net System Inventory = 820 - 223 = 597
    assert (tb_report.total_in - tb_report.total_out) == Decimal("597.0000")

    # Verify breakdown across physical stores:
    # WH001 (292 + 205) + WH002 (100) = 497 + 100 = 597
    wh001_total = store.get_balance(wh001.id, cement.id) + store.get_balance(
        wh001.id, rebar.id
    )
    wh002_total = store.get_balance(wh002.id, cement.id)
    assert wh001_total + wh002_total == Decimal("597.0000")

    # Verify reference and logistics tracking on report lines
    istv_line = next(i for i in tb_report.items if i.transaction_type == "ISTV")
    assert istv_line.plate_no == "3-A12345"
    assert istv_line.driver_name == "Abebe"

    srv_line = next(i for i in tb_report.items if i.transaction_type == "SRV")
    assert srv_line.in_quantity == Decimal("15.0000")

    # --------------------------------------------------------------------------
    # Step 10: Excel Export Verification
    # --------------------------------------------------------------------------
    excel_stream = await report_service.export_trial_balance_excel()
    assert excel_stream is not None

    wb = openpyxl.load_workbook(io.BytesIO(excel_stream.getvalue()))
    assert "Trial Balance" in wb.sheetnames
    ws = wb["Trial Balance"]

    # Header in row 4
    headers = [ws.cell(row=4, column=c).value for c in range(1, 19)]
    assert "Date" in headers
    assert "Tx Number" in headers
    assert "Type" in headers
    assert "Item Code" in headers
    assert "IN" in headers
    assert "OUT" in headers

    # 9 data rows (rows 5 through 13)
    data_rows = [
        ws.cell(row=r, column=2).value for r in range(5, 14)
    ]  # Tx Number column
    assert len(data_rows) == 9
    assert all(val is not None for val in data_rows)

    # Row 14: Total Row
    assert ws.cell(row=14, column=1).value == "TOTAL"
    assert ws.cell(row=14, column=14).value == 820.0
    assert ws.cell(row=14, column=15).value == 223.0


# ==============================================================================
# 4. HTTP API End-to-End Integration Route Tests
# ==============================================================================


@pytest.mark.asyncio
async def test_e2e_api_http_endpoints(e2e_system, token_factory):
    """
    Verifies that all 10 steps function end-to-end via real FastAPI HTTP requests,
    validating JWT authentication, request models, status codes, and streaming responses.
    """
    wh001: Warehouse = e2e_system["wh001"]
    wh002: Warehouse = e2e_system["wh002"]
    project: Project = e2e_system["project"]
    cement: Item = e2e_system["cement"]
    rebar: Item = e2e_system["rebar"]
    user: User = e2e_system["user"]
    tx_service: TransactionService = e2e_system["tx_service"]
    reconciliation_service: ReconciliationService = e2e_system["reconciliation_service"]
    report_service: ReportService = e2e_system["report_service"]

    # Configure FastAPI dependency overrides
    app.dependency_overrides[get_transaction_service] = lambda: tx_service
    app.dependency_overrides[get_reconciliation_service] = lambda: reconciliation_service
    app.dependency_overrides[get_report_service] = lambda: report_service

    token = token_factory(role="ADMIN", user_id=str(user.id))
    headers = {"Authorization": f"Bearer {token}"}

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        # Step 1: POST /transactions/grv
        grv_res = await client.post(
            "/api/v1/transactions/grv",
            json={
                "warehouse_id": str(wh001.id),
                "project_id": str(project.id),
                "supplier_name": "National Cement Share Co.",
                "lines": [
                    {"item_id": str(cement.id), "quantity": 500, "unit": "BAG"},
                    {"item_id": str(rebar.id), "quantity": 200, "unit": "PCS"},
                ],
            },
            headers=headers,
        )
        assert grv_res.status_code == 201
        grv_data = grv_res.json()
        assert grv_data["transaction_type"] == "GRV"

        # Step 2: POST /transactions/siv
        siv_res = await client.post(
            "/api/v1/transactions/siv",
            json={
                "warehouse_id": str(wh001.id),
                "project_id": str(project.id),
                "requested_from": "Civil Works Section",
                "issued_by_name": "Almaz Storekeeper",
                "checked_by_name": "Dawit Inspector",
                "received_by_name": "Chala Site Supervisor",
                "approved_by_name": "Bekele Project Manager",
                "lines": [
                    {"item_id": str(cement.id), "quantity": 120, "unit": "BAG"}
                ],
            },
            headers=headers,
        )
        assert siv_res.status_code == 201
        siv_data = siv_res.json()
        assert siv_data["transaction_type"] == "SIV"
        siv_id = siv_data["id"]

        # Step 3: POST /transactions/istv
        istv_res = await client.post(
            "/api/v1/transactions/istv",
            json={
                "source_warehouse_id": str(wh001.id),
                "destination_warehouse_id": str(wh002.id),
                "plate_no": "3-A12345",
                "driver_name": "Abebe",
                "lines": [
                    {"item_id": str(cement.id), "quantity": 100, "unit": "BAG"}
                ],
            },
            headers=headers,
        )
        assert istv_res.status_code == 201
        istv_data = istv_res.json()
        assert istv_data["transaction_type"] == "ISTV"
        istv_id = istv_data["id"]

        # Step 4: POST /transactions/istrv (Partial receipt 60)
        istrv1_res = await client.post(
            "/api/v1/transactions/istrv",
            json={
                "istv_transaction_id": istv_id,
                "destination_warehouse_id": str(wh002.id),
                "lines": [{"item_id": str(cement.id), "quantity": 60}],
            },
            headers=headers,
        )
        assert istrv1_res.status_code == 201

        # Step 5: POST /transactions/istrv (Final receipt 40)
        istrv2_res = await client.post(
            "/api/v1/transactions/istrv",
            json={
                "istv_transaction_id": istv_id,
                "destination_warehouse_id": str(wh002.id),
                "lines": [{"item_id": str(cement.id), "quantity": 40}],
            },
            headers=headers,
        )
        assert istrv2_res.status_code == 201

        # Step 6: POST /transactions/srv (Return 15 referencing SIV)
        srv_res = await client.post(
            "/api/v1/transactions/srv",
            json={
                "reference_siv_id": siv_id,
                "warehouse_id": str(wh001.id),
                "lines": [
                    {"item_id": str(cement.id), "quantity": 15, "unit": "BAG"}
                ],
            },
            headers=headers,
        )
        assert srv_res.status_code == 201

        # Step 7: POST /reconciliation/commit
        rec_res = await client.post(
            "/api/v1/reconciliation/commit",
            json={
                "warehouse_id": str(wh001.id),
                "count_reference": "Q3-HTTP-AUDIT-WH001",
                "adjustment_reason": "Quarterly Physical Audit Variance",
                "counts": [
                    {"item_id": str(cement.id), "counted_quantity": 292},
                    {"item_id": str(rebar.id), "counted_quantity": 205},
                ],
            },
            headers=headers,
        )
        assert rec_res.status_code == 200
        rec_data = rec_res.json()
        assert rec_data["adjusted_lines_count"] == 2

        # Step 9: GET /reports/trial-balance
        tb_res = await client.get("/api/v1/reports/trial-balance", headers=headers)
        assert tb_res.status_code == 200
        tb_data = tb_res.json()
        assert tb_data["total_count"] == 9
        assert Decimal(str(tb_data["total_in"])) == Decimal("820.0000")
        assert Decimal(str(tb_data["total_out"])) == Decimal("223.0000")

        # Step 10: GET /reports/trial-balance/export
        export_res = await client.get(
            "/api/v1/reports/trial-balance/export", headers=headers
        )
        assert export_res.status_code == 200
        assert (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            in export_res.headers["content-type"]
        )
        wb = openpyxl.load_workbook(io.BytesIO(export_res.content))
        assert "Trial Balance" in wb.sheetnames

    # Clean up overrides
    app.dependency_overrides.clear()


# ==============================================================================
# 5. Boundary & Security Invariant Tests
# ==============================================================================


@pytest.mark.asyncio
async def test_e2e_reconciliation_zero_variance(e2e_system):
    """Zero variance physical count posts no compensating adjustments."""
    wh001 = e2e_system["wh001"]
    cement = e2e_system["cement"]
    user = e2e_system["user"]
    tx_service = e2e_system["tx_service"]
    reconciliation_service = e2e_system["reconciliation_service"]

    # Initial stock: 100 bags
    await tx_service.create_and_post_grv(
        user_id=user.id,
        payload=GRVCreate(
            warehouse_id=wh001.id,
            lines=[
                TransactionLineCreate(
                    item_id=cement.id, quantity=Decimal("100.0000"), unit="BAG"
                )
            ],
        ),
    )

    # Physical count matches 100
    res = await reconciliation_service.commit_reconciliation_adjustments(
        user_id=user.id,
        payload=ReconciliationCommitRequest(
            warehouse_id=wh001.id,
            count_reference="MATCH-COUNT",
            adjustment_reason="Routine physical audit",
            counts=[
                PhysicalCountItemInput(
                    item_id=cement.id, counted_quantity=Decimal("100.0000")
                )
            ],
        ),
    )
    assert "No adjustments needed" in res["message"]
    assert res["adjusted_lines_count"] == 0
    assert res["transaction_number"] is None


@pytest.mark.asyncio
async def test_e2e_prevent_downward_adjustment_exceeding_stock(e2e_system):
    """ADR-002: Reconciliations cannot reduce stock below zero."""
    wh001 = e2e_system["wh001"]
    cement = e2e_system["cement"]
    user = e2e_system["user"]
    reconciliation_service = e2e_system["reconciliation_service"]

    # Stock is 0, attempting count of -10 is blocked by schema and service
    with pytest.raises(Exception):
        await reconciliation_service.commit_reconciliation_adjustments(
            user_id=user.id,
            payload=ReconciliationCommitRequest(
                warehouse_id=wh001.id,
                count_reference="ILLEGAL-COUNT",
                adjustment_reason="Negative audit",
                counts=[
                    PhysicalCountItemInput(
                        item_id=cement.id, counted_quantity=Decimal("-10.0000")
                    )
                ],
            ),
        )


@pytest.mark.asyncio
async def test_e2e_prevent_over_receiving_transfer(e2e_system):
    """ADR-005: Over-receiving on an in-transit transfer is strictly blocked."""
    wh001 = e2e_system["wh001"]
    wh002 = e2e_system["wh002"]
    cement = e2e_system["cement"]
    user = e2e_system["user"]
    tx_service = e2e_system["tx_service"]

    # Setup: 50 in WH001
    await tx_service.create_and_post_grv(
        user_id=user.id,
        payload=GRVCreate(
            warehouse_id=wh001.id,
            lines=[
                TransactionLineCreate(
                    item_id=cement.id, quantity=Decimal("50.0000"), unit="BAG"
                )
            ],
        ),
    )
    # Transfer 50
    istv_tx = await tx_service.create_and_post_istv(
        user_id=user.id,
        payload=ISTVCreate(
            source_warehouse_id=wh001.id,
            destination_warehouse_id=wh002.id,
            lines=[
                TransactionLineCreate(
                    item_id=cement.id, quantity=Decimal("50.0000"), unit="BAG"
                )
            ],
        ),
    )

    # Attempt to receive 60 (remaining is 50)
    with pytest.raises(HTTPException) as exc:
        await tx_service.create_and_post_istrv(
            user_id=user.id,
            payload=ISTRVCreate(
                istv_transaction_id=istv_tx.id,
                destination_warehouse_id=wh002.id,
                lines=[ISTRVLineCreate(item_id=cement.id, quantity=Decimal("60.0000"))],
            ),
        )
    assert exc.value.status_code == 400
    assert "over-receiving is prohibited" in exc.value.detail.lower()
