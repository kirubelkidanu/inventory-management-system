# backend/tests/test_reports.py
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
import io
from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

from httpx import ASGITransport, AsyncClient
from jose import jwt
import openpyxl
import pytest

from app.api.routes.reports import get_report_service
from app.core.config import settings
from app.main import app
from app.schemas.report import (
    InTransitReportItem,
    StockBalanceItem,
    StockBalanceResponse,
    TrialBalanceItem,
    TrialBalanceResponse,
)
from app.services.report import ReportService


@pytest.fixture
def token_factory():
    def make_token(role: str = "VIEWER", user_id: str = "user-123") -> str:
        return jwt.encode(
            {
                "sub": user_id,
                "email": f"{user_id}@example.com",
                "role": role,
                "exp": datetime.now(timezone.utc) + timedelta(minutes=30),
            },
            settings.JWT_SECRET,
            algorithm=settings.JWT_ALGORITHM,
        )

    return make_token


class FakeReportService:
    async def get_trial_balance_report(
        self,
        date_from=None,
        date_to=None,
        warehouse_id=None,
        project_id=None,
        item_id=None,
        category_id=None,
        transaction_type=None,
        search=None,
        page=1,
        page_size=50,
    ):
        item = TrialBalanceItem(
            movement_id=uuid4(),
            movement_date=date.today(),
            created_at=datetime.now(timezone.utc),
            transaction_id=uuid4(),
            transaction_number="GRV-20260918-0001",
            transaction_type="GRV",
            reference_number="INV-999",
            item_id=uuid4(),
            item_code="01-CM-00001",
            item_description="Portland Cement 42.5N",
            category_name="CEMENT",
            unit="BAG",
            warehouse_id=uuid4(),
            warehouse_code="WH001",
            warehouse_name="Main Warehouse",
            project_name="Kazanchis",
            plate_no="3-A12345",
            driver_name="Abebe Bikila",
            in_quantity=Decimal("100.0000"),
            out_quantity=Decimal("0.0000"),
            signed_quantity=Decimal("100.0000"),
            running_balance=Decimal("100.0000"),
            status="POSTED",
            entered_by_name="Store Clerk",
        )
        return TrialBalanceResponse(
            total_count=1,
            page=page,
            page_size=page_size,
            items=[item],
            total_in=Decimal("100.0000"),
            total_out=Decimal("0.0000"),
        )

    async def export_trial_balance_excel(
        self,
        date_from=None,
        date_to=None,
        warehouse_id=None,
        project_id=None,
        item_id=None,
        category_id=None,
        transaction_type=None,
        search=None,
    ):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Trial Balance"
        ws.append(["Date", "Tx Number", "IN", "OUT"])
        ws.append(["2026-09-18", "GRV-001", 100.0, 0.0])
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    async def get_stock_balance_report(
        self,
        warehouse_id=None,
        category_id=None,
        search=None,
    ):
        item = StockBalanceItem(
            warehouse_id=uuid4(),
            warehouse_code="WH001",
            warehouse_name="Main Warehouse",
            project_name="Kazanchis",
            category_name="CEMENT",
            item_id=uuid4(),
            item_code="01-CM-00001",
            item_description="Portland Cement 42.5N",
            unit="BAG",
            quantity_on_hand=Decimal("100.0000"),
            quantity_reserved=Decimal("20.0000"),
            quantity_available=Decimal("80.0000"),
        )
        return StockBalanceResponse(total_count=1, items=[item])

    async def get_in_transit_report(
        self,
        source_warehouse_id=None,
        destination_warehouse_id=None,
        status=None,
    ):
        item = InTransitReportItem(
            transfer_record_id=uuid4(),
            istv_transaction_id=uuid4(),
            istv_number="ISTV-20260918-0001",
            transaction_date=date.today(),
            source_warehouse_name="WH001",
            destination_warehouse_name="WH002",
            plate_no="3-B54321",
            driver_name="Tariku Tesfaye",
            status="IN_TRANSIT",
            item_id=uuid4(),
            item_code="01-CM-00001",
            item_description="Portland Cement",
            unit="BAG",
            sent_quantity=Decimal("50.0000"),
            received_quantity=Decimal("10.0000"),
            remaining_quantity=Decimal("40.0000"),
        )
        return [item]


# ===========================================================================
# 1. Service Layer Tests
# ===========================================================================


@pytest.mark.asyncio
async def test_service_trial_balance_report():
    mock_session = AsyncMock()
    service = ReportService(mock_session)

    wh_id = uuid4()
    item_id = uuid4()
    tx_id = uuid4()
    mov_id1 = uuid4()
    mov_id2 = uuid4()

    fake_rows = [
        {
            "movement_id": mov_id1,
            "movement_date": date(2026, 9, 1),
            "created_at": datetime(2026, 9, 1, 10, 0, tzinfo=timezone.utc),
            "transaction_id": tx_id,
            "transaction_number": "GRV-001",
            "transaction_type": "GRV",
            "invoice_no": "INV-101",
            "received_grv_no": None,
            "siv_no": None,
            "requested_no": None,
            "istv_no": None,
            "external_reference": None,
            "item_id": item_id,
            "item_code": "01-CM-00001",
            "item_description": "Cement 42.5N",
            "category_name": "CEMENT",
            "unit": "BAG",
            "warehouse_id": wh_id,
            "warehouse_code": "WH001",
            "warehouse_name": "Main Warehouse",
            "project_name": "Kazanchis",
            "plate_no": None,
            "driver_name": None,
            "movement_type": "IN",
            "quantity": Decimal("150.0000"),
            "signed_quantity": Decimal("150.0000"),
            "running_balance": Decimal("150.0000"),
            "status": "POSTED",
            "entered_by_name": "Admin User",
        },
        {
            "movement_id": mov_id2,
            "movement_date": date(2026, 9, 2),
            "created_at": datetime(2026, 9, 2, 11, 0, tzinfo=timezone.utc),
            "transaction_id": uuid4(),
            "transaction_number": "SIV-001",
            "transaction_type": "SIV",
            "invoice_no": None,
            "received_grv_no": None,
            "siv_no": "SIV-202",
            "requested_no": None,
            "istv_no": None,
            "external_reference": None,
            "item_id": item_id,
            "item_code": "01-CM-00001",
            "item_description": "Cement 42.5N",
            "category_name": "CEMENT",
            "unit": "BAG",
            "warehouse_id": wh_id,
            "warehouse_code": "WH001",
            "warehouse_name": "Main Warehouse",
            "project_name": "Kazanchis",
            "plate_no": None,
            "driver_name": None,
            "movement_type": "OUT",
            "quantity": Decimal("40.0000"),
            "signed_quantity": Decimal("-40.0000"),
            "running_balance": Decimal("110.0000"),
            "status": "POSTED",
            "entered_by_name": "Store Keeper",
        },
    ]

    service.repo.get_trial_balance = AsyncMock(
        return_value=(fake_rows, 2, Decimal("150.0000"), Decimal("40.0000"))
    )

    response = await service.get_trial_balance_report(
        warehouse_id=wh_id, page=1, page_size=50
    )

    assert response.total_count == 2
    assert response.total_in == Decimal("150.0000")
    assert response.total_out == Decimal("40.0000")
    assert len(response.items) == 2

    # Verify IN row mapping
    in_row = response.items[0]
    assert in_row.in_quantity == Decimal("150.0000")
    assert in_row.out_quantity == Decimal("0.0000")
    assert in_row.reference_number == "INV-101"

    # Verify OUT row mapping
    out_row = response.items[1]
    assert out_row.in_quantity == Decimal("0.0000")
    assert out_row.out_quantity == Decimal("40.0000")
    assert out_row.signed_quantity == Decimal("-40.0000")
    assert out_row.reference_number == "SIV-202"


@pytest.mark.asyncio
async def test_service_export_trial_balance_excel():
    mock_session = AsyncMock()
    service = ReportService(mock_session)

    raw_rows = [
        {
            "movement_id": uuid4(),
            "movement_date": date(2026, 9, 1),
            "created_at": datetime(2026, 9, 1, 10, 0, tzinfo=timezone.utc),
            "transaction_id": uuid4(),
            "transaction_number": "GRV-001",
            "transaction_type": "GRV",
            "invoice_no": "INV-101",
            "received_grv_no": None,
            "siv_no": None,
            "requested_no": None,
            "istv_no": None,
            "external_reference": None,
            "item_id": uuid4(),
            "item_code": "01-CM-00001",
            "item_description": "Cement",
            "category_name": "CEMENT",
            "unit": "BAG",
            "warehouse_id": uuid4(),
            "warehouse_code": "WH001",
            "warehouse_name": "Main Warehouse",
            "project_name": "Kazanchis",
            "plate_no": None,
            "driver_name": None,
            "movement_type": "IN",
            "quantity": Decimal("100.0000"),
            "signed_quantity": Decimal("100.0000"),
            "running_balance": Decimal("100.0000"),
            "status": "POSTED",
            "entered_by_name": "Admin",
        }
    ]
    service.repo.get_all_trial_balance_movements = AsyncMock(return_value=raw_rows)

    excel_stream = await service.export_trial_balance_excel()
    assert isinstance(excel_stream, io.BytesIO)

    # Validate workbook structure
    wb = openpyxl.load_workbook(excel_stream, data_only=True)
    assert "Trial Balance" in wb.sheetnames
    ws = wb["Trial Balance"]
    assert "DETAILED INVENTORY MOVEMENT" in str(ws["A1"].value)
    # Check headers at row 4
    headers = [cell.value for cell in ws[4]]
    assert "Tx Number" in headers
    assert "IN" in headers
    assert "OUT" in headers


@pytest.mark.asyncio
async def test_service_stock_balance_available_computation():
    mock_session = AsyncMock()
    service = ReportService(mock_session)

    raw_balances = [
        {
            "warehouse_id": uuid4(),
            "warehouse_code": "WH001",
            "warehouse_name": "Main",
            "project_name": "Kazanchis",
            "category_name": "STEEL",
            "item_id": uuid4(),
            "item_code": "02-ST-00001",
            "item_description": "Rebar 12mm",
            "unit": "PCS",
            "quantity_on_hand": Decimal("500.0000"),
            "quantity_reserved": Decimal("50.0000"),
            "quantity_available": Decimal("450.0000"),
        }
    ]
    service.repo.get_current_stock_balances = AsyncMock(return_value=raw_balances)

    result = await service.get_stock_balance_report()
    assert result.total_count == 1
    assert result.items[0].quantity_available == Decimal("450.0000")
    assert result.items[0].quantity_on_hand == Decimal("500.0000")


@pytest.mark.asyncio
async def test_service_in_transit_report():
    mock_session = AsyncMock()
    service = ReportService(mock_session)

    raw_transfers = [
        {
            "transfer_record_id": uuid4(),
            "istv_transaction_id": uuid4(),
            "istv_number": "ISTV-001",
            "transaction_date": date.today(),
            "source_warehouse_name": "WH001",
            "destination_warehouse_name": "WH002",
            "plate_no": "3-A12345",
            "driver_name": "Driver A",
            "status": "IN_TRANSIT",
            "item_id": uuid4(),
            "item_code": "01-CM-00001",
            "item_description": "Cement",
            "unit": "BAG",
            "sent_quantity": Decimal("100.0000"),
            "received_quantity": Decimal("30.0000"),
            "remaining_quantity": Decimal("70.0000"),
        }
    ]
    service.repo.get_in_transit_transfers = AsyncMock(return_value=raw_transfers)

    result = await service.get_in_transit_report()
    assert len(result) == 1
    assert result[0].remaining_quantity == Decimal("70.0000")
    assert result[0].istv_number == "ISTV-001"


# ===========================================================================
# 2. API Endpoints & RBAC Tests
# ===========================================================================


@pytest.mark.asyncio
async def test_api_trial_balance_success(token_factory):
    fake_service = FakeReportService()
    app.dependency_overrides[get_report_service] = lambda: fake_service
    token = token_factory(role="VIEWER")

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get(
            "/api/v1/reports/trial-balance",
            headers={"Authorization": f"Bearer {token}"},
        )

    app.dependency_overrides.clear()
    assert response.status_code == 200
    data = response.json()
    assert data["total_count"] == 1
    assert len(data["items"]) == 1
    assert Decimal(str(data["total_in"])) == Decimal("100.0000")
    assert data["items"][0]["transaction_type"] == "GRV"


@pytest.mark.asyncio
async def test_api_trial_balance_export_success(token_factory):
    fake_service = FakeReportService()
    app.dependency_overrides[get_report_service] = lambda: fake_service
    token = token_factory(role="ADMIN")

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get(
            "/api/v1/reports/trial-balance/export",
            headers={"Authorization": f"Bearer {token}"},
        )

    app.dependency_overrides.clear()
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment; filename=" in response.headers["content-disposition"]


@pytest.mark.asyncio
async def test_api_stock_balances_success(token_factory):
    fake_service = FakeReportService()
    app.dependency_overrides[get_report_service] = lambda: fake_service
    token = token_factory(role="STORE_KEEPER")

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get(
            "/api/v1/reports/stock-balances",
            headers={"Authorization": f"Bearer {token}"},
        )

    app.dependency_overrides.clear()
    assert response.status_code == 200
    data = response.json()
    assert data["total_count"] == 1
    assert Decimal(str(data["items"][0]["quantity_available"])) == Decimal("80.0000")


@pytest.mark.asyncio
async def test_api_in_transit_success(token_factory):
    fake_service = FakeReportService()
    app.dependency_overrides[get_report_service] = lambda: fake_service
    token = token_factory(role="INVENTORY_MANAGER")

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get(
            "/api/v1/reports/in-transit",
            headers={"Authorization": f"Bearer {token}"},
        )

    app.dependency_overrides.clear()
    assert response.status_code == 200
    data = response.json()
    assert len(data) == 1
    assert Decimal(str(data[0]["remaining_quantity"])) == Decimal("40.0000")


@pytest.mark.asyncio
@pytest.mark.parametrize("role", ["ADMIN", "INVENTORY_MANAGER", "STORE_KEEPER", "VIEWER"])
async def test_all_authenticated_roles_can_access_reports(token_factory, role):
    fake_service = FakeReportService()
    app.dependency_overrides[get_report_service] = lambda: fake_service
    token = token_factory(role=role)

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        res1 = await client.get(
            "/api/v1/reports/trial-balance",
            headers={"Authorization": f"Bearer {token}"},
        )
        res2 = await client.get(
            "/api/v1/reports/stock-balances",
            headers={"Authorization": f"Bearer {token}"},
        )
        res3 = await client.get(
            "/api/v1/reports/in-transit",
            headers={"Authorization": f"Bearer {token}"},
        )

    app.dependency_overrides.clear()
    assert res1.status_code == 200
    assert res2.status_code == 200
    assert res3.status_code == 200


@pytest.mark.asyncio
async def test_reports_unauthenticated_rejected():
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        res1 = await client.get("/api/v1/reports/trial-balance")
        res2 = await client.get("/api/v1/reports/trial-balance/export")
        res3 = await client.get("/api/v1/reports/stock-balances")
        res4 = await client.get("/api/v1/reports/in-transit")

    assert res1.status_code in (401, 403)
    assert res2.status_code in (401, 403)
    assert res3.status_code in (401, 403)
    assert res4.status_code in (401, 403)
